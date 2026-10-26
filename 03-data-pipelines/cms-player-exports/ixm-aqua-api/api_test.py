"""
Grassfish Player Export to Excel
=================================
Fetches all players from Grassfish IXM Aqua API and exports to a formatted Excel file.

Steps:
  1. Authenticate with Grassfish (OAuth client_credentials)
  2. Fetch all players with status, information, location, identification
  3. Map to a pandas DataFrame and export as .xlsx

Usage:
  1. Create a .env file with GRASSFISH_CLIENT_ID, GRASSFISH_CLIENT_SECRET, GRASSFISH_REGION
  2. pip install -r requirements.txt
  3. python grassfish_to_excel.py
"""

import os
import json
import logging
from datetime import datetime, timezone

from dotenv import load_dotenv
import requests

load_dotenv()  # reads .env file in the same directory
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────────────────────────────────────
# CONFIG — edit these or set as environment variables
# ──────────────────────────────────────────────────────────────────────────────

GRASSFISH_CLIENT_ID = os.getenv("GRASSFISH_CLIENT_ID", "YOUR_CLIENT_ID")
GRASSFISH_CLIENT_SECRET = os.getenv("GRASSFISH_CLIENT_SECRET", "YOUR_CLIENT_SECRET")
GRASSFISH_REGION = os.getenv("GRASSFISH_REGION", "as01")  # as01 / eu01 / us01

# Output file — use a timestamp so you never overwrite previous exports
OUTPUT_DIR = os.getenv("OUTPUT_DIR", ".")
OUTPUT_FILENAME = os.getenv("OUTPUT_FILENAME", "")  # leave blank for auto-timestamped name

# ──────────────────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger(__name__)


# ──────────────────────────────────────────────────────────────────────────────
# STEP 1 — Get Grassfish OAuth Token
# ──────────────────────────────────────────────────────────────────────────────

def get_grassfish_token(
    client_id: str = GRASSFISH_CLIENT_ID,
    client_secret: str = GRASSFISH_CLIENT_SECRET,
    region: str = GRASSFISH_REGION,
) -> str:
    identity_url = f"https://identity.{region}.signageplayer.app/connect/token"
    resp = requests.post(
        identity_url,
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "spconnect",
        },
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    resp.raise_for_status()
    data = resp.json()

    access_token = data.get("access_token")
    if not access_token:
        raise RuntimeError(f"Token request failed: {json.dumps(data)}")

    expires_in = data.get("expires_in", "?")
    log.info("Token acquired. Expires in %ss.", expires_in)
    return access_token


# ──────────────────────────────────────────────────────────────────────────────
# STEP 2 — Fetch All Players from Grassfish
# ──────────────────────────────────────────────────────────────────────────────

def fetch_players(access_token: str, region: str = GRASSFISH_REGION) -> list[dict]:
    api_base = f"https://connect.{region}.signageplayer.app/api/v1"
    resp = requests.post(
        f"{api_base}/selectplayers",
        json={},
        headers={
            "Authorization": f"Bearer {access_token}",
            "Content-Type": "application/json",
            "x-request-status": "true",
            "x-request-information": "true",
            "x-request-location": "true",
            "x-request-identification": "true",
        },
    )
    resp.raise_for_status()
    data = resp.json()

    success = data.get("success", {})
    if not success.get("ok"):
        raise RuntimeError(f"API returned unsuccessful response: {json.dumps(success)}")

    players = data.get("players", [])
    total_count = data.get("totalCount", len(players))
    log.info("Fetched %d of %d players.", len(players), total_count)
    return players


# ──────────────────────────────────────────────────────────────────────────────
# STEP 3 — Map players to rows and export as Excel
# ──────────────────────────────────────────────────────────────────────────────

def map_player_to_row(player: dict) -> dict:
    info = player.get("information") or {}
    status = player.get("status") or {}
    location = player.get("location") or {}

    is_online_raw = info.get("isOnline", False)
    is_online = is_online_raw in ("True", True)

    return {
        "Player ID":              str(player.get("id", "")),
        "Player Name":            player.get("name", ""),
        "Player Code":            player.get("code", ""),
        "Location Name":          location.get("name", ""),
        "Location ID":            str(location.get("id", "")),
        "Location Code":          location.get("code", ""),
        "Player Application":     player.get("playerApplication", ""),
        "Status Calculated":      status.get("statusCalculated", ""),
        "Status From Player":     status.get("statusFromPlayer", ""),
        "User Defined Status":    status.get("userDefinedStatus", ""),
        "Forgiving Status":       status.get("forgivingStatus", ""),
        "Is Online":              is_online,
        "Is Error":               status.get("isError", False),
        "Is Not Started":         status.get("isNotStarted", False),
        "Is Empty Playlist":      status.get("isEmptyPlaylist", False),
        "Serial Number":          info.get("computerSerialNumber", ""),
        "MAC Address":            info.get("macAddress", ""),
        "IP Address":             info.get("ipAddress", ""),
        "Network Type":           info.get("networkConnectionType", ""),
        "Network Strength":       info.get("networkStrength", ""),
        "OS Version":             info.get("osVersion", ""),
        "Player Version":         info.get("playerVersion", ""),
        "Computer Manufacturer":  info.get("computerManufacturer", ""),
        "Computer Name":          info.get("computerName", ""),
        "CPU Model":              info.get("cpuModel", ""),
        "Screen Serial":          info.get("screenOneSerialNumber", ""),
        "Screen Model":           info.get("screenOneModel", ""),
        "Screen Firmware":        info.get("screenOneFirmware", ""),
        "Last Communication":     info.get("lastCommunicationDateTime", ""),
        "Last Synced":            datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
        "Tags":                   ", ".join(
            (t.get("name") if isinstance(t, dict) else str(t))
            for t in (player.get("tags") or [])
        ),
    }


def export_to_excel(players: list[dict], output_path: str) -> str:
    rows = [map_player_to_row(p) for p in players]
    df = pd.DataFrame(rows)

    # Sort by Location Name then Player Name for readability
    df.sort_values(by=["Location Name", "Player Name"], inplace=True, ignore_index=True)

    # Write to Excel with openpyxl engine for formatting
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Players")
        ws = writer.sheets["Players"]

        # ── Header styling ──
        header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="2D5B8A")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # ── Auto-fit column widths ──
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = len(str(col_name))
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)

        # ── Data font ──
        data_font = Font(name="Arial", size=10)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.font = data_font

        # ── Freeze the header row ──
        ws.freeze_panes = "A2"

        # ── Auto-filter ──
        ws.auto_filter.ref = ws.dimensions

    log.info("Exported %d players to %s", len(df), output_path)
    return output_path


# ──────────────────────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────────────────────

def main():
    log.info("=== Grassfish Player Export Started ===")

    # Step 1: Authenticate
    access_token = get_grassfish_token()

    # Step 2: Fetch players
    players = fetch_players(access_token)

    if not players:
        log.warning("No players returned — nothing to export.")
        return

    # Step 3: Export to Excel
    if OUTPUT_FILENAME:
        filename = OUTPUT_FILENAME
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"grassfish_players_{timestamp}.xlsx"

    output_path = os.path.join(OUTPUT_DIR, filename)
    export_to_excel(players, output_path)

    log.info("=== Done. File: %s ===", output_path)


if __name__ == "__main__":
    main()