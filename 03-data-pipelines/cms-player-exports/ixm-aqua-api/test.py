"""
Grassfish → Excel Player Export (All Clients)
===============================================
Fetches players across all sub-accounts and exports to a formatted .xlsx
matching the CMS DP Information CSV format.

Usage:
  1. Create a .env with GRASSFISH_CLIENT_ID, GRASSFISH_CLIENT_SECRET, GRASSFISH_REGION
  2. pip install -r requirements.txt
  3. python grassfish_to_excel.py
"""

import os
import json
import logging
import time
from datetime import datetime, timezone

from dotenv import load_dotenv
import requests
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

load_dotenv()

# ── Config ────────────────────────────────────────────────────────────────────
GRASSFISH_CLIENT_ID = os.getenv("GRASSFISH_CLIENT_ID", "YOUR_CLIENT_ID")
GRASSFISH_CLIENT_SECRET = os.getenv("GRASSFISH_CLIENT_SECRET", "YOUR_CLIENT_SECRET")
GRASSFISH_REGION = os.getenv("GRASSFISH_REGION", "as01")
OUTPUT_DIR = os.getenv("OUTPUT_DIR", ".")
OUTPUT_FILENAME = os.getenv("OUTPUT_FILENAME", "")

IDENTITY_URL = f"https://identity.{GRASSFISH_REGION}.signageplayer.app/connect/token"
API_BASE = f"https://connect.{GRASSFISH_REGION}.signageplayer.app/api"
RATE_LIMIT_DELAY = 0.25  # 250ms between client requests

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-8s  %(message)s", datefmt="%Y-%m-%d %H:%M:%S")
log = logging.getLogger(__name__)


# ── Step 1: OAuth Token ──────────────────────────────────────────────────────

def get_token() -> str:
    resp = requests.post(IDENTITY_URL, data={
        "grant_type": "client_credentials",
        "client_id": GRASSFISH_CLIENT_ID,
        "client_secret": GRASSFISH_CLIENT_SECRET,
        "scope": "spconnect",
    }, headers={"Content-Type": "application/x-www-form-urlencoded"})
    resp.raise_for_status()
    data = resp.json()
    token = data.get("access_token")
    if not token:
        raise RuntimeError(f"Token request failed: {json.dumps(data)}")
    log.info("Token acquired. Expires in %ss.", data.get("expires_in", "?"))
    return token


# ── Step 2: Fetch all clients ────────────────────────────────────────────────

def fetch_clients(token: str) -> list[dict]:
    resp = requests.post(f"{API_BASE}/selectclients", json={}, headers={
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
    }, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    clients = data.get("clients", [])
    log.info("Found %d clients.", len(clients))
    return clients


# ── Step 3: Fetch players per client ─────────────────────────────────────────

def fetch_players_for_client(token: str, client_id: str) -> list[dict]:
    resp = requests.post(f"{API_BASE}/selectplayers", json={}, headers={
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
        "x-clientId": client_id,
        "x-request-status": "true",
        "x-request-information": "true",
        "x-request-location": "true",
        "x-request-identification": "true",
    }, timeout=60)
    resp.raise_for_status()
    data = resp.json()
    return data.get("players", [])


def fetch_all_players(token: str, clients: list[dict]) -> list[dict]:
    all_players = []
    for i, client in enumerate(clients, start=1):
        cid = client.get("id", "")
        cname = client.get("name", "")
        try:
            players = fetch_players_for_client(token, cid)
            for p in players:
                p["_client_name"] = cname
                p["_client_id"] = cid
            all_players.extend(players)
            log.info("[%d/%d] %s — %d players", i, len(clients), cname, len(players))
        except Exception as e:
            log.error("[%d/%d] %s — FAILED: %s", i, len(clients), cname, e)

        if i < len(clients):
            time.sleep(RATE_LIMIT_DELAY)

    log.info("Total: %d players across %d clients.", len(all_players), len(clients))
    return all_players


# ── Step 4: Map to CSV-matching columns and export ───────────────────────────

def derive_screen_status(player: dict) -> str:
    status = player.get("status") or {}
    calculated = status.get("statusCalculated", "")
    if calculated == "OK":
        return "Online"
    from_player = status.get("statusFromPlayer", "")
    if from_player == "OK":
        return "Online"
    return "Offline"


def map_player_to_row(player: dict) -> dict:
    info = player.get("information") or {}
    status = player.get("status") or {}
    location = player.get("location") or {}
    ident = player.get("identification") or {}

    screen_models = info.get("screenModels") or []
    if not screen_models and info.get("screenOneModel"):
        screen_models = [info["screenOneModel"]]

    return {
        "Client":            player.get("_client_name", ""),
        "PlayerName":        player.get("name", ""),
        "PlayerId":          player.get("id", ""),
        "Location":          location.get("name", ""),
        "LocationId":        location.get("id", ""),
        "MountedDirection":  "",  # not available via API
        "PlayerFormat":      "",  # not available via API
        "AutoPair":          str(ident.get("isAllowAutoIdentification", False)).upper(),
        "PlayerCode":        player.get("code", ""),
        "Tags":              ", ".join(
            (t.get("name") if isinstance(t, dict) else str(t))
            for t in (player.get("tags") or [])
        ),
        "ScreenModels":      ", ".join(screen_models),
        "FirstContactDate":  info.get("insertedDateTime", ""),
        "LastContactDate":   info.get("lastCommunicationDateTime", "")
                             or status.get("lastCommunicationDateTime", ""),
        "SerialNumber":      info.get("computerSerialNumber", ""),
        "ScreenStatus":      derive_screen_status(player),
        "ClientId":          player.get("_client_id", ""),
    }


def export_to_excel(players: list[dict], output_path: str) -> str:
    rows = [map_player_to_row(p) for p in players]
    df = pd.DataFrame(rows)

    df.sort_values(by=["Client", "PlayerName"], inplace=True, ignore_index=True)

    timestamp = datetime.now(timezone.utc).strftime("%a %d %b %Y %H:%M")
    df[f"LAST UPDATED: {timestamp}"] = ""

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="CMS DP Information")
        ws = writer.sheets["CMS DP Information"]

        header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="2D5B8A")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = len(str(col_name))
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)

        data_font = Font(name="Arial", size=10)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.font = data_font

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    log.info("Exported %d players to %s", len(df), output_path)
    return output_path


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    log.info("=== Grassfish Player Export Started ===")

    token = get_token()
    clients = fetch_clients(token)

    if not clients:
        log.warning("No clients found — nothing to export.")
        return

    players = fetch_all_players(token, clients)

    if not players:
        log.warning("No players returned — nothing to export.")
        return

    if OUTPUT_FILENAME:
        filename = OUTPUT_FILENAME
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"CMS_DP_Information_{timestamp}.xlsx"

    output_path = os.path.join(OUTPUT_DIR, filename)
    export_to_excel(players, output_path)
    log.info("=== Done. File: %s ===", output_path)


if __name__ == "__main__":
    main()