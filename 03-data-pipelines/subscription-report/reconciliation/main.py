import os
#!/usr/bin/env python3
"""
Digital Signage Subscription Report Generator
Generates monthly subscription timeline per customer group from Airtable.
Cross-references start dates against the sales xlsx and highlights discrepancies.

Usage:
    python3 generate_subscription_report.py [path_to_sales_xlsx]

Dependencies:
    pip install requests openpyxl rapidfuzz python-dateutil
"""

import sys
import time
import json
from pathlib import Path
from datetime import date, datetime
from dateutil.relativedelta import relativedelta

import requests
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz, process as rfprocess

TOKEN = os.environ["AT_PAT"]  # was hard-coded; rotated at handoff
HEADERS = {"Authorization": f"Bearer {TOKEN}"}

SUBS_BASE  = "appRor1ipbYAb4K7z"
SUBS_TABLE = "tblyNA73D9jldeEir"
JOBS_BASE  = "appdo1WwONYvmYMWg"
JOBS_TABLE = "tblwHj9cO37wp1fVx"

# Location lookup tables (records that the link/lookup fields point to)
SUB_LOC_BASE  = "appRor1ipbYAb4K7z"   # same base as subscriptions
SUB_LOC_TABLE = "tblp7sqsv25DVYs7T"   # Inventory table (contains location records)
JOB_LOC_BASE  = "appdo1WwONYvmYMWg"   # same base as jobs
JOB_LOC_TABLE = "tbl5Nj3zs3nxKPdEv"   # Job locations table

# Subscription fields (with returnFieldsByFieldId=true)
S_LOC    = "fldHUCgQFipY5UaR4"  # Location - Installation (lookup -> record IDs)
S_START  = "fld95WxbXtbHfktSU"  # Start date
S_END    = "fldSO7ONdpFqVJnte"  # End date (formula -> array)
S_STATUS = "fld22c3r9LiQw6j4F"  # Active / Expired
S_CGROUP = "fldcFDSqwac3Vo4qb"  # Customer group string

# Job fields (with returnFieldsByFieldId=true)
J_LOC       = "flddTDwVWOq7fc1Je"  # Location (link -> record IDs)
J_COMPLETED = "fldu0PmyaUa3gAMjd"  # Date completed
J_SCHEDULED = "fld19en5UMabiUnqG"  # Date scheduled (fallback)
J_TYPE      = "fldk5E1PkIy1ZrM1e"  # Job type (plain string with returnFieldsByFieldId)
J_STATUS    = "fld7q5wjIg47nltes"  # Job status (plain string)

SKIP_SHEETS = {
    "Copperwired ", "Copperwired TH", "Template", "Summary PMC",
    "Copperwired", "Inventory-Patlitas Grid",
}

START_MONTH = date(2022, 4, 1)
END_MONTH   = date.today().replace(day=1)

FILL_HEADER      = PatternFill("solid", fgColor="1F3864")
FILL_TOTAL_ROW   = PatternFill("solid", fgColor="BDD7EE")
FILL_ACTIVE      = PatternFill("solid", fgColor="C6EFCE")
FILL_DISCREPANCY = PatternFill("solid", fgColor="FFEB9C")

FONT_HEADER = Font(color="FFFFFF", bold=True, size=9, name="Arial")
FONT_BODY   = Font(size=9, name="Arial")
FONT_BOLD   = Font(size=9, bold=True, name="Arial")
FONT_TOTAL  = Font(size=9, bold=True, name="Arial")

SALES_XLSX  = sys.argv[1] if len(sys.argv) > 1 else "new_digital_signage_subscription_tracker_sales.xlsx"
OUTPUT_PATH = "subscription_report_output.xlsx"


def get_months():
    months, m = [], START_MONTH
    while m <= END_MONTH:
        months.append(m)
        m += relativedelta(months=1)
    return months

def month_label(d):
    return d.strftime("%b %Y")

def parse_date(s):
    if not s:
        return None
    try:
        return date.fromisoformat(str(s)[:10])
    except Exception:
        return None

def is_active(start_str, end_str, month_date):
    s = parse_date(start_str)
    e = parse_date(end_str)
    if not s or not e:
        return False
    month_end = month_date + relativedelta(months=1) - relativedelta(days=1)
    return s <= month_end and e >= month_date

def fmt_date(d):
    return d.strftime("%d-%b-%Y") if d else ""

def make_sheet_name(cgroup):
    parts = cgroup.split(" - ", 1)
    name = parts[1] if len(parts) > 1 else cgroup
    for ch in r'\/*?:[]':
        name = name.replace(ch, "")
    return name[:31].strip()


def fetch_all(base, table, fields=None, use_field_ids=True):
    url = f"https://api.airtable.com/v0/{base}/{table}"
    records, offset = [], None
    page = 0
    while True:
        params = [("pageSize", "100")]
        if use_field_ids:
            params.append(("returnFieldsByFieldId", "true"))
        if fields:
            params += [("fields[]", f) for f in fields]
        if offset:
            params.append(("offset", offset))
        r = requests.get(url, headers=HEADERS, params=params)
        if r.status_code == 429:
            print("  Rate limited, waiting 30s...")
            time.sleep(30)
            continue
        if r.status_code != 200:
            print(f"  Error {r.status_code}: {r.text[:200]}")
            break
        data = r.json()
        records.extend(data.get("records", []))
        page += 1
        print(f"  Page {page}: {len(records)} records...", end="\r", flush=True)
        offset = data.get("offset")
        if not offset:
            break
        time.sleep(0.22)
    print(f"  Done: {len(records)} records total        ")
    return records


def build_loc_map(base, table, label="locations"):
    """
    Fetch all records from a table and build {record_id: primary_field_value}.
    Fetches WITHOUT returnFieldsByFieldId so we get field names as keys.
    The primary field for location tables is "ID" (full path like
    "THA - Apple - Advice - iStore @ Phichit").
    """
    print(f"  Fetching {label} from {base}/{table}...")
    records = fetch_all(base, table, fields=None, use_field_ids=False)
    loc_map = {}
    for rec in records:
        rid = rec["id"]
        fields = rec.get("fields", {})
        # Try "ID" field first (primary field in Inventory / Job locations tables)
        name = fields.get("ID")
        if not name:
            # Fallback: find any string field with the location path pattern
            for v in fields.values():
                if isinstance(v, str) and v.count(" - ") >= 2:
                    name = v
                    break
        if name:
            loc_map[rid] = str(name)
    print(f"  {len(loc_map)} {label} mapped")
    return loc_map


def load_subscriptions(loc_map):
    print("\n[2/5] Fetching subscriptions...")
    records = fetch_all(SUBS_BASE, SUBS_TABLE, [S_LOC, S_START, S_END, S_STATUS, S_CGROUP])
    subs = []
    skipped = 0
    for rec in records:
        f = rec.get("fields", {})

        # Location: lookup field returns array of record IDs
        loc_ids = f.get(S_LOC)
        loc = None
        if isinstance(loc_ids, list) and loc_ids:
            # Each element is a plain record ID string
            loc = loc_map.get(loc_ids[0])

        # End date: formula field returns array
        end_raw = f.get(S_END)
        end = end_raw[0] if isinstance(end_raw, list) and end_raw else end_raw

        cgroup = f.get(S_CGROUP, "")

        if loc and cgroup:
            subs.append({
                "location": loc,
                "start":    f.get(S_START),
                "end":      end,
                "status":   f.get(S_STATUS, ""),
                "cgroup":   cgroup,
            })
        else:
            skipped += 1

    print(f"  {len(subs)} valid records ({skipped} skipped)")
    return subs


def load_jobs(loc_map):
    print("\n[3/5] Fetching installation jobs...")
    records = fetch_all(JOBS_BASE, JOBS_TABLE, [J_LOC, J_COMPLETED, J_SCHEDULED, J_TYPE, J_STATUS])
    jobs_by_loc = {}
    for rec in records:
        f = rec.get("fields", {})

        # Job type and status: plain strings with returnFieldsByFieldId
        jtype = f.get(J_TYPE, "")
        jstat = f.get(J_STATUS, "")

        if jtype != "Installation":
            continue
        if jstat not in ("Done", "Installed"):
            continue

        # Location: link field returns array of record ID strings
        loc_ids = f.get(J_LOC)
        loc = None
        if isinstance(loc_ids, list) and loc_ids:
            loc = loc_map.get(loc_ids[0])
        if not loc:
            continue

        completed = parse_date(f.get(J_COMPLETED))
        scheduled = parse_date(f.get(J_SCHEDULED))
        d = completed or scheduled
        if not d:
            continue

        if loc not in jobs_by_loc or d < jobs_by_loc[loc]["earliest"]:
            jobs_by_loc[loc] = {"earliest": d, "completed": completed}

    print(f"  {len(jobs_by_loc)} locations with installation jobs")
    return jobs_by_loc


def coerce_to_date(val):
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, (int, float)) and 20000 < val < 60000:
        try:
            return (datetime(1899, 12, 30) + relativedelta(days=int(val))).date()
        except Exception:
            return None
    if isinstance(val, str):
        val = val.strip()
        if not val:
            return None
        for fmt in ("%d-%b-%Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%B %d, %Y", "%b %d, %Y"):
            try:
                return datetime.strptime(val, fmt).date()
            except ValueError:
                pass
        try:
            from dateutil.parser import parse as dp
            return dp(val).date()
        except Exception:
            pass
    return None


def find_col_idx(headers, candidates):
    hl = [str(h).lower().strip() if h else "" for h in headers]
    for c in candidates:
        for i, h in enumerate(hl):
            if c.lower() in h:
                return i
    return None


def load_xlsx_dates(path):
    p = Path(path)
    if not p.exists():
        print(f"  Warning: sales xlsx not found at '{path}'")
        return {}
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    except Exception as e:
        print(f"  Warning: could not open xlsx: {e}")
        return {}

    xlsx_dates = {}
    sheets_processed = 0

    for sname in wb.sheetnames:
        if sname.strip() in SKIP_SHEETS or sname in SKIP_SHEETS:
            continue
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hrow_idx, headers = None, None
        for i, row in enumerate(rows[:5]):
            if row and any("location" in str(c).lower() for c in row if c):
                hrow_idx = i
                headers = list(row)
                break
        if hrow_idx is None:
            continue
        loc_col = find_col_idx(headers, ["location name", "location - name", "store name"])
        if loc_col is None:
            continue
        start_col = find_col_idx(headers, [
            "(free) start date", "free) start date", "free start date",
            "1st start date", "1st year start", "contract start date",
            "ma start date",
        ])
        if start_col is None:
            continue
        sheets_processed += 1
        for row in rows[hrow_idx + 1:]:
            if not row or len(row) <= max(loc_col, start_col):
                continue
            loc_val   = row[loc_col]
            start_val = row[start_col]
            if not loc_val or not start_val:
                continue
            loc_str = str(loc_val).strip()
            if not loc_str or loc_str.lower() in ("location name", "store name", "none"):
                continue
            d = coerce_to_date(start_val)
            if d and (loc_str not in xlsx_dates or d < xlsx_dates[loc_str]):
                xlsx_dates[loc_str] = d

    wb.close()
    print(f"  {sheets_processed} sheets parsed, {len(xlsx_dates)} unique locations found")
    return xlsx_dates


def build_short_name_index(at_locations):
    index = {}
    for full in at_locations:
        short = full.rsplit(" - ", 1)[-1].strip()
        if short not in index:
            index[short] = full
        else:
            index[full] = full
    return index


def fuzzy_find(xlsx_loc, short_index, threshold=75):
    result = rfprocess.extractOne(
        xlsx_loc, short_index.keys(),
        scorer=fuzz.token_sort_ratio, score_cutoff=threshold,
    )
    if result:
        matched_key, score, _ = result
        return short_index[matched_key], score
    return None, 0


def build_report(subs, jobs_by_loc, xlsx_dates):
    months = get_months()

    by_cgroup = {}
    for s in subs:
        by_cgroup.setdefault(s["cgroup"], {}).setdefault(s["location"], []).append(s)

    all_at_locs = {loc for locs in by_cgroup.values() for loc in locs}
    short_index = build_short_name_index(all_at_locs)

    discrepancy_locs = set()
    uncertain_matches = []

    for xlsx_loc, xlsx_start in xlsx_dates.items():
        at_full, score = fuzzy_find(xlsx_loc, short_index)
        if not at_full:
            continue
        if score < 90:
            uncertain_matches.append((xlsx_loc, at_full, score))
        for cg, locs in by_cgroup.items():
            if at_full in locs:
                at_starts = [parse_date(s["start"]) for s in locs[at_full] if s["start"]]
                if not at_starts:
                    continue
                at_earliest = min(at_starts)
                if at_earliest.replace(day=1) != xlsx_start.replace(day=1):
                    discrepancy_locs.add(at_full)
                break

    if uncertain_matches:
        print(f"\n  {len(uncertain_matches)} fuzzy matches with score < 90:")
        for x, a, sc in uncertain_matches[:15]:
            print(f"     '{x}' -> '{a}' (score {sc})")
        if len(uncertain_matches) > 15:
            print(f"     ... and {len(uncertain_matches) - 15} more")

    print(f"\n  {len(discrepancy_locs)} locations with start-date discrepancy vs xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for cg in sorted(by_cgroup.keys()):
        sname = make_sheet_name(cg)
        ws = wb.create_sheet(title=sname)

        month_labels = [month_label(m) for m in months]
        headers = ["Location", "Date - Completed (Job)", "Installation Date"] + month_labels
        n_fixed = 3
        n_months = len(months)

        ws.append(headers)
        for col_i in range(1, len(headers) + 1):
            c = ws.cell(1, col_i)
            c.fill = FILL_HEADER
            c.font = FONT_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

        locs = by_cgroup[cg]
        for loc_name in sorted(locs.keys()):
            loc_subs = locs[loc_name]
            job_info = jobs_by_loc.get(loc_name, {})
            completed_date = job_info.get("completed")
            install_date   = job_info.get("earliest")
            is_disc        = loc_name in discrepancy_locs

            month_counts = [
                sum(1 for s in loc_subs if is_active(s["start"], s["end"], m))
                for m in months
            ]
            row = [loc_name, fmt_date(completed_date), fmt_date(install_date)] + month_counts
            ws.append(row)
            ri = ws.max_row

            ws.cell(ri, 1).font = FONT_BOLD if is_disc else FONT_BODY
            ws.cell(ri, 2).font = FONT_BODY
            ws.cell(ri, 3).font = FONT_BODY
            ws.cell(ri, 2).alignment = Alignment(horizontal="center")
            ws.cell(ri, 3).alignment = Alignment(horizontal="center")

            fill = FILL_DISCREPANCY if is_disc else FILL_ACTIVE
            for mi, cnt in enumerate(month_counts):
                c = ws.cell(ri, n_fixed + 1 + mi)
                c.alignment = Alignment(horizontal="center")
                c.font = FONT_BODY
                if cnt > 0:
                    c.fill = fill

        total_row_idx = ws.max_row + 1
        ws.cell(total_row_idx, 1).value = "TOTAL"
        data_start_row = 2
        data_end_row   = total_row_idx - 1
        for mi in range(n_months):
            col_letter = get_column_letter(n_fixed + 1 + mi)
            c = ws.cell(total_row_idx, n_fixed + 1 + mi)
            if data_end_row >= data_start_row:
                c.value = f"=SUM({col_letter}{data_start_row}:{col_letter}{data_end_row})"
            else:
                c.value = 0
            c.font = FONT_TOTAL
            c.fill = FILL_TOTAL_ROW
            c.alignment = Alignment(horizontal="center")
        for col_i in range(1, n_fixed + 1):
            c = ws.cell(total_row_idx, col_i)
            c.font = FONT_TOTAL
            c.fill = FILL_TOTAL_ROW

        legend_row = total_row_idx + 2
        for i, (lfill, label) in enumerate([
            (FILL_ACTIVE,      "Active subscription"),
            (FILL_DISCREPANCY, "Start date differs from sales xlsx"),
        ]):
            r = legend_row + i
            ws.cell(r, 1).fill  = lfill
            ws.cell(r, 2).value = label
            ws.cell(r, 2).font  = FONT_BODY

        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 20
        for mi in range(n_months):
            ws.column_dimensions[get_column_letter(n_fixed + 1 + mi)].width = 9
        ws.freeze_panes = "D2"

    # Safety: if no sheets were created, add a placeholder
    if not wb.sheetnames:
        ws = wb.create_sheet("No Data")
        ws["A1"] = "No subscription data found"
        print("  WARNING: No customer groups found, created empty placeholder sheet")

    return wb


def main():
    print("=" * 60)
    print("  Digital Signage Subscription Report Generator")
    print("=" * 60)
    print(f"  Month range : {month_label(START_MONTH)} -> {month_label(END_MONTH)}")
    print(f"  Sales xlsx  : {SALES_XLSX}")
    print(f"  Output      : {OUTPUT_PATH}")
    print("=" * 60)

    # Step 1: Build location name maps (record_id -> full location name)
    print("\n[1/5] Building location name maps...")
    sub_loc_map = build_loc_map(SUB_LOC_BASE, SUB_LOC_TABLE, "subscription locations")
    job_loc_map = build_loc_map(JOB_LOC_BASE, JOB_LOC_TABLE, "job locations")

    # Step 2: Fetch subscriptions (resolve location IDs via sub_loc_map)
    subs = load_subscriptions(sub_loc_map)

    # Step 3: Fetch jobs (resolve location IDs via job_loc_map)
    jobs_by_loc = load_jobs(job_loc_map)

    # Step 4: Cross-reference sales xlsx
    print(f"\n[4/5] Loading sales xlsx for cross-reference...")
    xlsx_dates = load_xlsx_dates(SALES_XLSX)

    # Step 5: Build report
    print(f"\n[5/5] Building report...")
    wb = build_report(subs, jobs_by_loc, xlsx_dates)

    wb.save(OUTPUT_PATH)
    print(f"\n  Saved: {OUTPUT_PATH}")
    print(f"  Sheets: {len(wb.sheetnames)}")
    print(f"  Months: {len(get_months())}")


if __name__ == "__main__":
    main()