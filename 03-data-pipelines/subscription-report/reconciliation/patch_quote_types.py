"""
Patch Quote Type + Quote Service in HubSpot
============================================
For each quote:
  1. PATCHes quote_type from the grouped xlsx
  2. Looks up the associated deal's "service" property
  3. Maps it to the quote's "quote_service" enum and PATCHes that too

Usage:
  export HS_API_KEY="your_token"
  python patch_quote_types.py

Requirements:
  pip install requests openpyxl
"""

import os, sys, time
import requests
from openpyxl import load_workbook

HS_API_KEY = os.environ.get("HS_API_KEY", "YOUR_PRIVATE_APP_TOKEN_HERE")
BASE_URL = "https://api.hubapi.com"
HEADERS = {"Authorization": f"Bearer {HS_API_KEY}", "Content-Type": "application/json"}
INPUT_FILE = "quote_types_grouped.xlsx"

# Mapping: deal "service" value → quote "quote_service" enum value
DEAL_SERVICE_TO_QUOTE_SERVICE = {
    "Digital Signage": "true",           # quote_service enum value for Digital Signage
    "Music Streaming": "music_streaming",
}


def get_associated_deal_id(quote_id: str) -> str | None:
    """Get the first associated deal ID for a quote. Tries multiple approaches."""
    for obj_type in ["deal", "deals"]:
        url = f"{BASE_URL}/crm/v4/objects/quotes/{quote_id}/associations/{obj_type}"
        resp = requests.get(url, headers=HEADERS)
        if resp.status_code == 200:
            results = resp.json().get("results", [])
            if results:
                return str(results[0]["toObjectId"])

    url = f"{BASE_URL}/crm/v3/objects/quotes/{quote_id}?associations=deal"
    resp = requests.get(url, headers=HEADERS)
    if resp.status_code == 200:
        assoc = resp.json().get("associations", {})
        for key in ["deals", "deal"]:
            results = assoc.get(key, {}).get("results", [])
            if results:
                return str(results[0]["id"])

    return None


def get_deal_service(deal_id: str) -> str | None:
    """Read the 'service' property from a deal and map it to quote_service enum value."""
    url = f"{BASE_URL}/crm/v3/objects/deals/{deal_id}?properties=service"
    resp = requests.get(url, headers=HEADERS)
    if resp.status_code != 200:
        return None
    deal_service = (resp.json().get("properties", {}).get("service") or "").strip()
    if not deal_service:
        return None
    return DEAL_SERVICE_TO_QUOTE_SERVICE.get(deal_service)


def patch_quote(record_id: str, properties: dict, deal_id: str | None = None) -> tuple[bool, str]:
    """
    PATCH quote properties. If simple PATCH fails with MISSING_QUOTE_ASSOCIATED_OBJECT,
    retry with deal association in request body, then batch update.
    """
    url = f"{BASE_URL}/crm/v3/objects/quotes/{record_id}"

    # Attempt 1: simple PATCH
    resp = requests.patch(url, headers=HEADERS, json={"properties": properties})
    if resp.status_code == 200:
        return True, ""

    if resp.status_code == 400 and "MISSING_QUOTE_ASSOCIATED_OBJECT" in resp.text:
        if not deal_id:
            deal_id = get_associated_deal_id(record_id)
        if not deal_id:
            return False, "ACCEPTED CPQ quote — no associated deal found"

        # Attempt 2: PATCH with associations
        body = {
            "properties": properties,
            "associations": [
                {
                    "to": {"id": deal_id},
                    "types": [{"associationCategory": "HUBSPOT_DEFINED", "associationTypeId": 64}],
                }
            ],
        }
        resp2 = requests.patch(url, headers=HEADERS, json=body)
        if resp2.status_code == 200:
            return True, ""

        # Attempt 3: batch update
        resp3 = requests.post(
            f"{BASE_URL}/crm/v3/objects/quotes/batch/update",
            headers=HEADERS,
            json={"inputs": [{"id": record_id, "properties": properties}]},
        )
        if resp3.status_code == 200:
            return True, ""

        return False, f"All attempts failed (deal={deal_id}). Last: {resp3.status_code}: {resp3.text[:200]}"

    return False, f"{resp.status_code}: {resp.text[:200]}"


def main():
    if HS_API_KEY == "YOUR_PRIVATE_APP_TOKEN_HERE":
        print("ERROR: Set HS_API_KEY environment variable.")
        sys.exit(1)

    if not os.path.exists(INPUT_FILE):
        print(f"ERROR: {INPUT_FILE} not found in current directory.")
        sys.exit(1)

    wb = load_workbook(INPUT_FILE, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    total = len(rows)
    print(f"Loaded {total} quotes from {INPUT_FILE}\n")

    success, failed, skipped = 0, 0, 0
    svc_updated, svc_skipped = 0, 0
    errors = []

    for i, row in enumerate(rows, 1):
        record_id = str(row[0]).strip() if row[0] else ""
        quote_number = str(row[1]) if row[1] else ""
        quote_title = str(row[2]) if row[2] else ""
        quote_type = str(row[4]).strip() if row[4] else ""

        label = f"[{i}/{total}] {quote_number} — {quote_title}"

        if not record_id or record_id == "None":
            print(f"  ⚠ SKIP {label} (no record ID)")
            skipped += 1
            continue

        if not quote_type or quote_type == "None":
            print(f"  ⚠ SKIP {label} (no quote type)")
            skipped += 1
            continue

        # Look up deal for quote_service
        deal_id = get_associated_deal_id(record_id)
        quote_service = None
        if deal_id:
            quote_service = get_deal_service(deal_id)

        # Build properties to patch
        props = {"quote_type": quote_type}
        if quote_service:
            props["quote_service"] = quote_service

        ok, err = patch_quote(record_id, props, deal_id)
        if ok:
            svc_label = f" | service={quote_service}" if quote_service else " | service=(no deal/empty)"
            print(f"  ✓ {label} → type={quote_type}{svc_label}")
            success += 1
            if quote_service:
                svc_updated += 1
            else:
                svc_skipped += 1
        else:
            print(f"  ✗ {label} → FAILED: {err}")
            errors.append({"record_id": record_id, "quote_number": quote_number, "error": err})
            failed += 1

        time.sleep(0.05)

    print(f"\n{'='*60}")
    print(f"Done.")
    print(f"  Quote Type:    ✓ {success} updated | ✗ {failed} failed | ⚠ {skipped} skipped")
    print(f"  Quote Service: ✓ {svc_updated} updated | ⚠ {svc_skipped} no deal/service")

    if errors:
        print("\nFailed records:")
        for e in errors:
            print(f"  {e['quote_number']} (ID: {e['record_id']}): {e['error']}")


if __name__ == "__main__":
    main()