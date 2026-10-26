"""
Extract all HubSpot quote line items into an Excel file.

Usage:
  export HS_API_KEY="your_token"
  python extract_quote_line_items.py

Output: quote_line_items.xlsx
"""

import os, sys, requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HS_API_KEY = os.environ.get("HS_API_KEY", "YOUR_PRIVATE_APP_TOKEN_HERE")
BASE_URL = "https://api.hubapi.com"
HEADERS = {"Authorization": f"Bearer {HS_API_KEY}", "Content-Type": "application/json"}
OUTPUT_FILE = "quote_line_items.xlsx"


def fetch_all_quotes():
    all_quotes = []
    after = None
    while True:
        url = f"{BASE_URL}/crm/v3/objects/quotes?limit=100&properties=hs_title,hs_quote_number,quote_type,hs_status"
        if after:
            url += f"&after={after}"
        resp = requests.get(url, headers=HEADERS)
        resp.raise_for_status()
        data = resp.json()
        all_quotes.extend(data.get("results", []))
        after = data.get("paging", {}).get("next", {}).get("after")
        if not after:
            break
    return all_quotes


def get_line_item_ids(quote_id):
    url = f"{BASE_URL}/crm/v4/objects/quotes/{quote_id}/associations/line_items"
    resp = requests.get(url, headers=HEADERS)
    resp.raise_for_status()
    return [str(r["toObjectId"]) for r in resp.json().get("results", [])]


def batch_read_line_items(ids):
    if not ids:
        return []
    results = []
    for i in range(0, len(ids), 100):
        batch = ids[i:i+100]
        resp = requests.post(
            f"{BASE_URL}/crm/v3/objects/line_items/batch/read",
            headers=HEADERS,
            json={"inputs": [{"id": lid} for lid in batch], "properties": ["name", "quantity", "quote_category"]},
        )
        resp.raise_for_status()
        results.extend(resp.json().get("results", []))
    return results


def main():
    if HS_API_KEY == "YOUR_PRIVATE_APP_TOKEN_HERE":
        print("ERROR: Set HS_API_KEY environment variable.")
        sys.exit(1)

    print("Fetching all quotes...")
    quotes = fetch_all_quotes()
    print(f"  Found {len(quotes)} quotes")

    rows = []
    for i, q in enumerate(quotes):
        qid = q["id"]
        props = q.get("properties", {})
        title = props.get("hs_title", "")
        number = props.get("hs_quote_number", "")
        current_qt = props.get("quote_type", "") or ""
        status = props.get("hs_status", "")

        print(f"  [{i+1}/{len(quotes)}] {number} - {title}...", end=" ")

        li_ids = get_line_item_ids(qid)
        if not li_ids:
            print("no line items")
            rows.append({
                "record_id": qid,
                "quote_number": number,
                "quote_title": title,
                "quote_status": status,
                "line_item_id": "",
                "line_item": "(no line items)",
                "quantity": "",
                "existing_quote_category": "",
                "existing_quote_type": current_qt,
                "quote_type": "",
            })
            continue

        line_items = batch_read_line_items(li_ids)
        print(f"{len(line_items)} line items")

        for li in line_items:
            li_props = li.get("properties", {})
            rows.append({
                "record_id": qid,
                "quote_number": number,
                "quote_title": title,
                "quote_status": status,
                "line_item_id": li["id"],
                "line_item": li_props.get("name", ""),
                "quantity": li_props.get("quantity", ""),
                "existing_quote_category": li_props.get("quote_category", "") or "",
                "existing_quote_type": current_qt,
                "quote_type": "",
            })

    # Build xlsx
    print(f"\nGenerating {OUTPUT_FILE} with {len(rows)} rows...")
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote Line Items"

    headers = [
        "Record ID (Quote)",
        "Quote Number",
        "Quote Title",
        "Quote Status",
        "Line Item ID",
        "Line Item",
        "Quantity",
        "Existing Quote Category (from line item)",
        "Existing Quote Type (on quote)",
        "Quote Type (FILL THIS IN)",
    ]

    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Arial", size=10)
    fill_yellow = PatternFill("solid", fgColor="FFFF00")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for r, row in enumerate(rows, 2):
        ws.cell(row=r, column=1, value=row["record_id"]).font = cell_font
        ws.cell(row=r, column=2, value=row["quote_number"]).font = cell_font
        ws.cell(row=r, column=3, value=row["quote_title"]).font = cell_font
        ws.cell(row=r, column=4, value=row["quote_status"]).font = cell_font
        ws.cell(row=r, column=5, value=row["line_item_id"]).font = cell_font
        ws.cell(row=r, column=6, value=row["line_item"]).font = cell_font
        ws.cell(row=r, column=7, value=row["quantity"]).font = cell_font
        ws.cell(row=r, column=8, value=row["existing_quote_category"]).font = cell_font
        ws.cell(row=r, column=9, value=row["existing_quote_type"]).font = cell_font
        # Quote Type column - highlighted yellow for user to fill in
        qt_cell = ws.cell(row=r, column=10, value=row["quote_type"])
        qt_cell.font = cell_font
        qt_cell.fill = fill_yellow

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 40
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 30
    ws.column_dimensions["I"].width = 25
    ws.column_dimensions["J"].width = 28

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_FILE)
    print(f"✅ Done! Saved to {OUTPUT_FILE}")
    print(f"\nInstructions:")
    print(f"  1. Open {OUTPUT_FILE}")
    print(f"  2. Fill in column J ('Quote Type') for each row")
    print(f"     Valid values: Hardware, Software, Installation, Delivery, Maintenance")
    print(f"  3. Multiple types per quote are fine - they'll be combined per quote")
    print(f"  4. Feed the file back to Claude for processing")


if __name__ == "__main__":
    main()