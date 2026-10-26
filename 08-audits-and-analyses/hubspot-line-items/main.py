"""
HubSpot Quote Line Item Exporter
=================================
Fetches all quotes + their line items, deals, and buyer companies,
then writes a clean Excel file for price analysis.

Usage:
    pip install requests openpyxl
    export HUBSPOT_API_KEY=pat-na1-xxxxxxxxxxxx
    python hubspot_quote_export.py

Output:
    hubspot_quotes_export.xlsx
"""

import os
import sys
import time
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

API_KEY = os.environ.get("HUBSPOT_API_KEY", "")
BASE_URL = "https://api.hubapi.com"
HEADERS = {"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"}
OUTPUT_FILE = "hubspot_quotes_export.xlsx"

QUOTE_PROPS = ["hs_quote_number", "hs_status", "hs_title", "hs_expiration_date", "hs_currency"]
DEAL_PROPS = ["dealname"]
COMPANY_PROPS = ["name"]
LINE_ITEM_PROPS = ["name", "price", "quantity", "hs_sku", "amount"]

RETRY_LIMIT = 3
RETRY_DELAY = 2  # seconds


# ---------------------------------------------------------------------------
# API Helpers
# ---------------------------------------------------------------------------

def get(url, params=None):
    """GET with retry logic for rate limits."""
    for attempt in range(RETRY_LIMIT):
        resp = requests.get(url, headers=HEADERS, params=params, timeout=30)
        if resp.status_code == 429:
            print(f"  Rate limited. Waiting {RETRY_DELAY}s...")
            time.sleep(RETRY_DELAY)
            continue
        resp.raise_for_status()
        return resp.json()
    raise Exception(f"Failed after {RETRY_LIMIT} retries: {url}")


def post(url, payload):
    for attempt in range(RETRY_LIMIT):
        resp = requests.post(url, headers=HEADERS, json=payload, timeout=30)
        if resp.status_code == 429:
            print(f"  Rate limited. Waiting {RETRY_DELAY}s...")
            time.sleep(RETRY_DELAY)
            continue
        resp.raise_for_status()
        return resp.json()
    raise Exception(f"Failed after {RETRY_LIMIT} retries: {url}")


# ---------------------------------------------------------------------------
# Fetchers
# ---------------------------------------------------------------------------

def fetch_all_quotes():
    """Page through all quotes and return list of quote objects."""
    quotes = []
    url = f"{BASE_URL}/crm/v3/objects/quotes"
    params = {
        "limit": 100,
        "properties": ",".join(QUOTE_PROPS),
        "associations": "line_items,deals",
    }
    print("Fetching quotes...")
    while True:
        data = get(url, params)
        quotes.extend(data.get("results", []))
        paging = data.get("paging", {}).get("next", {})
        after = paging.get("after")
        if not after:
            break
        params["after"] = after
        print(f"  ...{len(quotes)} quotes so far")
    print(f"  Total quotes fetched: {len(quotes)}")
    return quotes


def fetch_object_by_id(object_type, object_id, properties):
    """Fetch a single CRM object by id."""
    url = f"{BASE_URL}/crm/v3/objects/{object_type}/{object_id}"
    params = {"properties": ",".join(properties)}
    try:
        data = get(url, params)
        return data.get("properties", {})
    except Exception as e:
        print(f"  Warning: could not fetch {object_type}/{object_id}: {e}")
        return {}


def fetch_associations(from_type, from_id, to_type):
    """Return list of associated object IDs."""
    url = f"{BASE_URL}/crm/v3/objects/{from_type}/{from_id}/associations/{to_type}"
    try:
        data = get(url)
        return [r["id"] for r in data.get("results", [])]
    except Exception:
        return []


def fetch_deal_company(deal_id):
    """Get the first associated company name for a deal."""
    company_ids = fetch_associations("deals", deal_id, "companies")
    if not company_ids:
        return ""
    props = fetch_object_by_id("companies", company_ids[0], COMPANY_PROPS)
    return props.get("name", "")


# ---------------------------------------------------------------------------
# Data Assembly
# ---------------------------------------------------------------------------

def extract_associations(quote, object_type):
    """Pull associated IDs from the embedded associations blob."""
    assoc_blob = quote.get("associations", {})
    type_data = assoc_blob.get(object_type, {}).get("results", [])
    return [item["id"] for item in type_data]


def process_quote(quote, index, total):
    props = quote.get("properties", {})
    quote_id = quote.get("id", "")
    quote_number = props.get("hs_quote_number", "")
    quote_status = props.get("hs_status", "")
    quote_name = props.get("hs_title", "")
    currency = props.get("hs_currency", "")

    print(f"  [{index}/{total}] Quote {quote_number or quote_id}")

    deal_ids = extract_associations(quote, "deals") or fetch_associations("quotes", quote_id, "deals")

    deal_name = ""
    company_name = ""
    if deal_ids:
        deal_props = fetch_object_by_id("deals", deal_ids[0], DEAL_PROPS)
        deal_name = deal_props.get("dealname", "")
        company_name = fetch_deal_company(deal_ids[0])

    line_item_ids = extract_associations(quote, "line_items") or fetch_associations("quotes", quote_id, "line_items")

    if not line_item_ids:
        return [{
            "quote_number": quote_number,
            "quote_status": quote_status,
            "quote_name": quote_name,
            "deal_name": deal_name,
            "company_name": company_name,
            "line_item": "(no line items)",
            "price": "",
            "quantity": "",
            "amount": "",
            "currency": currency,
        }]

    rows = []
    for li_id in line_item_ids:
        li_props = fetch_object_by_id("line_items", li_id, LINE_ITEM_PROPS)
        price_raw = li_props.get("price", "")
        qty_raw = li_props.get("quantity", "")
        amt_raw = li_props.get("amount", "")
        rows.append({
            "quote_number": quote_number,
            "quote_status": quote_status,
            "quote_name": quote_name,
            "deal_name": deal_name,
            "company_name": company_name,
            "line_item": li_props.get("name", ""),
            "price": float(price_raw) if price_raw else "",
            "quantity": float(qty_raw) if qty_raw else "",
            "amount": float(amt_raw) if amt_raw else "",
            "currency": currency,
        })
    return rows


def build_rows(quotes):
    from concurrent.futures import ThreadPoolExecutor, as_completed

    rows = []
    total = len(quotes)

    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = {
            executor.submit(process_quote, q, i, total): i
            for i, q in enumerate(quotes, 1)
        }
        for future in as_completed(futures):
            try:
                rows.extend(future.result())
            except Exception as e:
                print(f"  Error processing quote: {e}")

    rows.sort(key=lambda r: r.get("quote_number") or "")
    return rows


# ---------------------------------------------------------------------------
# Excel Writer
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=10)
ALT_FILL = PatternFill("solid", start_color="EBF3FB")
BORDER = Border(
    bottom=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


COLUMNS = [
    ("Quote Number",   "quote_number",   22, "str"),
    ("Quote Status",   "quote_status",   18, "str"),
    ("Quote Name",     "quote_name",     30, "str"),
    ("Deal Name",      "deal_name",      30, "str"),
    ("Buyer Company",  "company_name",   28, "str"),
    ("Line Item",      "line_item",      35, "str"),
    ("Currency",       "currency",       12, "str"),
    ("Unit Price",     "price",          16, "currency"),
    ("Qty",            "quantity",       10, "number"),
    ("Total Amount",   "amount",         18, "currency"),
]


def write_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote Line Items"
    ws.freeze_panes = "A2"

    # Header row
    for col_idx, (header, _, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, row in enumerate(rows, 2):
        alt = (row_idx % 2 == 0)
        for col_idx, (_, key, _, fmt) in enumerate(COLUMNS, 1):
            value = row.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if fmt == "str" else CENTER

            if alt and fmt == "str":
                cell.fill = ALT_FILL

            if fmt == "currency" and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00'
                if alt:
                    cell.fill = ALT_FILL
            elif fmt == "number" and isinstance(value, (int, float)):
                cell.number_format = '#,##0.##'
                if alt:
                    cell.fill = ALT_FILL

    # Summary sheet
    ws_sum = wb.create_sheet("Summary")
    ws_sum["A1"] = "Total Quotes"
    ws_sum["B1"] = len({r["quote_number"] for r in rows})
    ws_sum["A2"] = "Total Line Item Rows"
    ws_sum["B2"] = len(rows)
    ws_sum["A3"] = "Grand Total Amount"
    ws_sum["B3"] = sum(r["amount"] for r in rows if isinstance(r.get("amount"), float))
    for cell in [ws_sum["A1"], ws_sum["A2"], ws_sum["A3"]]:
        cell.font = Font(bold=True, name="Arial", size=10)
    ws_sum["B3"].number_format = '#,##0.00'

    wb.save(OUTPUT_FILE)
    print(f"\nSaved: {OUTPUT_FILE}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    if not API_KEY:
        print("Error: Set HUBSPOT_API_KEY environment variable first.")
        print("  export HUBSPOT_API_KEY=pat-na1-xxxxxxxxxxxx")
        sys.exit(1)

    print("=" * 55)
    print("  HubSpot Quote Exporter")
    print("=" * 55)

    quotes = fetch_all_quotes()

    print(f"\nBuilding rows for {len(quotes)} quotes...")
    rows = build_rows(quotes)

    print(f"\nWriting Excel file with {len(rows)} rows...")
    write_excel(rows)

    print("\nDone! Two sheets inside:")
    print(f"  - 'Quote Line Items'  ({len(rows)} rows)")
    print(f"  - 'Summary'           (totals at a glance)")


if __name__ == "__main__":
    main()